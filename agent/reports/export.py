"""Export a Markdown report to PDF (headless Chromium) or Word .docx (python-docx). RTL Arabic."""

from __future__ import annotations

import html as _html
import re
from pathlib import Path
from typing import List, Tuple

Block = Tuple  # ("h", level, text) | ("li", text) | ("p", text) | ("hr",)


def _is_table_row(ln: str) -> bool:
    return bool(re.match(r"^\s*\|.*\|\s*$", ln))


def _is_table_sep(ln: str) -> bool:
    return bool(re.match(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$", ln)) and "-" in ln


def _row_cells(ln: str) -> List[str]:
    return [c.strip() for c in ln.strip().strip("|").split("|")]


def parse_blocks(md: str) -> List[Block]:
    blocks: List[Block] = []
    lines = (md or "").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue
        # جدول Markdown: صف رؤوس + سطر فاصل + صفوف
        if _is_table_row(ln) and i + 1 < len(lines) and _is_table_sep(lines[i + 1]):
            header = _row_cells(ln)
            rows = []
            i += 2
            while i < len(lines) and _is_table_row(lines[i]):
                rows.append(_row_cells(lines[i])); i += 1
            blocks.append(("table", header, rows)); continue
        if re.match(r"^\s*(---|\*\*\*|___)\s*$", ln):
            blocks.append(("hr",)); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)", ln)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip())); i += 1; continue
        if re.match(r"^\s*[-•*]\s+", ln):
            blocks.append(("li", re.sub(r"^\s*[-•*]\s+", "", ln))); i += 1; continue
        blocks.append(("p", ln.strip())); i += 1
    return blocks


# ---------- inline ----------
_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
_BOLD = re.compile(r"\*\*([^*]+)\*\*")
_CODE = re.compile(r"`([^`]+)`")
_URL = re.compile(r"(?<![\"(>])(https?://[^\s<]+)")


def _inline_html(text: str) -> str:
    t = _html.escape(text)
    t = _LINK.sub(r'<a href="\2">\1</a>', t)
    t = _URL.sub(r'<a href="\1">\1</a>', t)
    t = _BOLD.sub(r"<strong>\1</strong>", t)
    t = _CODE.sub(r"<code>\1</code>", t)
    return t


def to_html_doc(md: str, title: str = "تقرير") -> str:
    blocks = parse_blocks(md)
    body, in_ul = [], False
    def close():
        nonlocal in_ul
        if in_ul: body.append("</ul>"); in_ul = False
    for b in blocks:
        if b[0] == "li":
            if not in_ul: body.append("<ul>"); in_ul = True
            body.append(f"<li>{_inline_html(b[1])}</li>")
        elif b[0] == "h":
            close(); lvl = min(b[1], 4); body.append(f"<h{lvl}>{_inline_html(b[2])}</h{lvl}>")
        elif b[0] == "hr":
            close(); body.append("<hr/>")
        elif b[0] == "table":
            close()
            _, header, rows = b
            th = "".join(f"<th>{_inline_html(c)}</th>" for c in header)
            trs = "".join("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in r) + "</tr>" for r in rows)
            body.append(f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>")
        else:
            close(); body.append(f"<p>{_inline_html(b[1])}</p>")
    close()
    return f"""<!DOCTYPE html><html lang="ar" dir="rtl"><head><meta charset="UTF-8">
<title>{_html.escape(title)}</title>
<style>
  @page {{ size: A4; margin: 20mm 16mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: "Noto Naskh Arabic", "Segoe UI", Tahoma, sans-serif; color: #2b2118; line-height: 1.9; font-size: 12pt; }}
  h1 {{ color: #a63f08; font-size: 20pt; border-bottom: 3px solid #a63f08; padding-bottom: 6px; margin: 0 0 14px; }}
  h2 {{ color: #1a5c53; font-size: 15pt; border-bottom: 1px solid #cdba93; padding-bottom: 4px; margin: 20px 0 8px; }}
  h3 {{ color: #40301c; font-size: 13pt; margin: 14px 0 6px; }}
  h4 {{ color: #5e4930; font-size: 12pt; margin: 10px 0 4px; }}
  p {{ margin: 6px 0; }}
  ul {{ margin: 6px 0; padding-inline-start: 22px; }}
  li {{ margin: 3px 0; }}
  a {{ color: #a63f08; text-decoration: none; }}
  strong {{ font-weight: 700; }}
  code {{ background: #f0e8d6; padding: 1px 5px; border-radius: 4px; font-size: 10pt; }}
  hr {{ border: 0; border-top: 1px solid #cdba93; margin: 14px 0; }}
  table {{ width: 100%; border-collapse: collapse; margin: 10px 0; font-size: 11pt; }}
  th, td {{ border: 1px solid #cdba93; padding: 6px 8px; text-align: right; }}
  th {{ background: #f0e8d6; color: #1a5c53; font-weight: 700; }}
  .footer {{ margin-top: 24px; padding-top: 8px; border-top: 1px solid #cdba93; color: #86704f; font-size: 9pt; }}
</style></head><body>
{''.join(body)}
<div class="footer">وكيل Odoo للمراقبة — تقرير مُصدَّر آلياً</div>
</body></html>"""


def to_pdf(md: str, out_path: Path, title: str = "تقرير") -> Path:
    from agent.tools.browser import get_renderer
    renderer = get_renderer()
    if not renderer.available():
        raise RuntimeError("محرك المتصفح غير مثبّت — لا يمكن إنشاء PDF")
    doc = to_html_doc(md, title)
    with renderer._lock:  # serialize browser use (small-RAM box)
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
            try:
                pg = b.new_page()
                pg.set_content(doc, wait_until="networkidle")
                pg.pdf(path=str(out_path), format="A4", print_background=True,
                       margin={"top": "18mm", "bottom": "18mm", "left": "14mm", "right": "14mm"})
            finally:
                b.close()
    return out_path


# ---------- DOCX ----------
def _set_rtl(paragraph) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def _add_runs(paragraph, text: str) -> None:
    """Split on **bold**; drop link markup to 'text (url)'."""
    text = _LINK.sub(r"\1 (\2)", text)
    for i, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith("**") else part)
        if part.startswith("**"):
            run.bold = True


def to_docx(md: str, out_path: Path, title: str = "تقرير") -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    doc.sections[0].right_to_left = True if hasattr(doc.sections[0], "right_to_left") else None

    for b in parse_blocks(md):
        if b[0] == "h":
            h = doc.add_heading(level=min(b[1], 4))
            h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(h)
            _add_runs(h, b[2])
        elif b[0] == "li":
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p)
            _add_runs(p, b[1])
        elif b[0] == "hr":
            doc.add_paragraph("—" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b[0] == "table":
            _, header, rows = b
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for j, c in enumerate(header):
                cell = t.rows[0].cells[j]; cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; _set_rtl(cell.paragraphs[0]); _add_runs(cell.paragraphs[0], c)
            for r in rows:
                cells = t.add_row().cells
                for j, c in enumerate(r[:len(header)]):
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; _set_rtl(cells[j].paragraphs[0]); _add_runs(cells[j].paragraphs[0], c)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p)
            _add_runs(p, b[1])
    doc.save(str(out_path))
    return out_path
